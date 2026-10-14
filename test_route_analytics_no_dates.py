"""Regression test: the website's Route Analytics has NO date filter — every trip (any year) counts, the From / To / Vehicle Type filters still
work, an old bookmarked URL that still carries date_from/date_to is simply ignored, and the Excel export covers the same trips. Runs against a
disposable copy of fleet.db with a few extra trips added; the real file is never modified.

Usage:  python3 test_route_analytics_no_dates.py
"""
import io
import os
import re
import shutil
import sqlite3
import sys
import tempfile

REPO = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, REPO)
os.chdir(REPO)
import requests                                                              # noqa: E402
def _blocked(*a, **k):
    raise RuntimeError('network blocked in tests')
requests.get = requests.post = _blocked
from flask import template_rendered                                           # noqa: E402
from openpyxl import load_workbook                                            # noqa: E402
import app as m                                                               # noqa: E402
from test_support import ensure_test_user                                     # noqa: E402

ok = fail = 0


def check(cond, msg):
    global ok, fail
    if cond:
        ok += 1
    else:
        fail += 1
        print(f'[FAIL] {msg}')


tmp = tempfile.mkdtemp(prefix='ranodates_')
dbp = os.path.join(tmp, 't.db')
src = sqlite3.connect('fleet.db'); src.backup(sqlite3.connect(dbp)); src.close()
uid = ensure_test_user(dbp)


def gdb():
    c = sqlite3.connect(dbp)
    c.row_factory = sqlite3.Row
    return c


m.get_db = gdb
m.app.config['TESTING'] = True
db = gdb()
# a route nobody else has, with trips in three different years — before the old default range, inside it, and after it
db.execute("DELETE FROM trips WHERE from_loc IN ('ZZTESTFROM','YYTESTFROM')")
veh = db.execute("SELECT id, type FROM vehicles ORDER BY id LIMIT 1").fetchone()
def trip(date, frm, to, rate, vtype='own', empty=0, lr=None):
    db.execute("INSERT INTO trips (date, lr_number, from_loc, to_loc, type, rate_type, rate, quantity, billed_amount, vehicle_id, is_empty, fuel_amount, driver_adv_amount) "
               "VALUES (?,?,?,?,?,'PER_MT',?,10,?,?,?,100,50)", (date, lr or f'T-{date}-{rate}', frm, to, vtype, rate, rate * 10, veh['id'], empty))
trip('2024-05-10', 'ZZTESTFROM', 'ZZTESTTO', 500)          # old
trip('2026-06-10', 'ZZTESTFROM', 'ZZTESTTO', 700)          # inside the old FY 2026-27 default
trip('2028-01-05', 'ZZTESTFROM', 'ZZTESTTO', 900)          # later than the old default
trip('2025-02-02', 'YYTESTFROM', 'YYTESTTO', 300, vtype='hired')
trip('2026-06-11', 'ZZTESTFROM', 'ZZTESTTO', 99999, empty=1, lr='EMPTYONE')   # an empty run must stay out
db.commit()

cl = m.app.test_client()
with cl.session_transaction() as s:
    s['user_id'] = uid
seen = []
def grab(sender, template, context, **kw):
    seen.append(context)
template_rendered.connect(grab, m.app)


def page(url):
    seen.clear()
    r = cl.get(url)
    ctx = next((c for c in seen if 'total_trips_sel' in c), None)
    return r, ctx


try:
    r, ctx = page('/route-analytics?tab=best-routes&from=ZZTESTFROM&to=ZZTESTTO')
    check(r.status_code == 200 and ctx is not None, 'page opens')
    check(ctx['total_trips_sel'] == 3, f"all three trips count, whatever their year ({ctx['total_trips_sel']})")
    html = r.data.decode()
    check('name="date_from"' not in html and 'name="date_to"' not in html and 'From Date' not in html and 'To Date' not in html, 'the filter bar has no date fields')
    check(all(x in html for x in ('name="from"', 'name="to"', 'name="vehicle_type"')), 'the From / To / Vehicle Type filters are still there')
    check('Selected Period' not in html, 'no leftover "Selected Period" wording on a page that has no period')

    r2, ctx2 = page('/route-analytics?tab=best-routes&from=ZZTESTFROM&to=ZZTESTTO&date_from=2026-04-01&date_to=2027-03-31')
    check(ctx2['total_trips_sel'] == 3, 'an old bookmarked URL carrying dates is ignored (still all 3 trips)')
    r3, ctx3 = page('/route-analytics?tab=best-routes&from=ZZTESTFROM&to=ZZTESTTO&date_from=2030-01-01&date_to=2030-12-31')
    check(ctx3['total_trips_sel'] == 3, 'even a date range with no trips in it changes nothing')

    # Route Rates: highest / lowest / average now come from every trip
    r, ctx = page('/route-analytics?tab=route-rates&from=ZZTESTFROM&to=ZZTESTTO')
    rows = ctx['rr_rows']
    check(len(rows) == 1 and rows[0]['trips'] == 3, f'Route Rates groups all three trips ({[(x["trips"]) for x in rows]})')
    check(rows[0]['highest'] == 900 and rows[0]['lowest'] == 500 and round(rows[0]['average']) == 700, f"rates span every year: highest/avg/lowest = {rows[0]['highest']}/{round(rows[0]['average'])}/{rows[0]['lowest']}")
    check(99999 not in (rows[0]['highest'],) and ctx['total_trips_sel'] == 3, 'the empty run stays excluded')
    check({p['month'] for p in ctx['trend_points']} >= {'2024-05', '2026-06', '2028-01'}, f"the Rate Trend chart covers every month with data ({[p['month'] for p in ctx['trend_points']][:6]}...)")

    # layout of the Route Rates tab: the table comes first, the Driver Advances / Fuel Taken cards sit below it
    rr_html = cl.get('/route-analytics?tab=route-rates&from=ZZTESTFROM&to=ZZTESTTO').data.decode()
    panel = rr_html[rr_html.index('id="ra-route-rates"'):]
    pos = {k: panel.find(k) for k in ('Highest Rate', 'class="rr-table"', 'Driver Advances (₹)', 'Fuel Taken (₹)</div>', 'Rate Trend (Average')}
    check(all(v >= 0 for v in pos.values()), 'Route Rates tab still shows the rate cards, the table, both advance cards and the charts')
    check(pos['Highest Rate'] < pos['class="rr-table"'] < pos['Driver Advances (₹)'] < pos['Fuel Taken (₹)</div>'] < pos['Rate Trend (Average'],
          'Route Rates tab order: rate cards, table, then Driver Advances and Fuel Taken, then the charts')
    check(panel.count('stat-compound-card') == 2 and '₹' in panel[pos['Driver Advances (₹)']:pos['Rate Trend (Average']], 'both advance cards still carry their numbers')

    # the other filters still narrow things
    r, ctx = page('/route-analytics?tab=best-routes&from=YYTESTFROM')
    check(ctx['total_trips_sel'] == 1, 'filtering by From still works')
    r, ctx = page('/route-analytics?tab=best-routes&from=ZZTESTFROM&to=ZZTESTTO&vehicle_type=hired')
    check(ctx['total_trips_sel'] == 0, 'filtering by Vehicle Type still works')
    r, ctx = page('/route-analytics?tab=best-routes')
    all_trips = db.execute("SELECT COUNT(*) FROM trips WHERE NOT " + m._EMPTY_RUN_SQL).fetchone()[0]
    check(ctx['total_trips_sel'] == all_trips, f"with no filters every non-empty trip in the database counts ({ctx['total_trips_sel']} of {all_trips})")

    # export: same scope
    r = cl.get('/route-analytics/export?from=ZZTESTFROM&to=ZZTESTTO')
    check(r.status_code == 200 and r.data[:2] == b'PK', 'Excel export downloads')
    wb = load_workbook(io.BytesIO(r.data))
    cells = [str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None]
    check(any('ZZTESTFROM' in c or 'ZZTESTTO' in c for c in cells), 'the export contains the route')
    trips_cell = [row for ws in wb for row in ws.iter_rows(values_only=True) if row and any(isinstance(v, str) and 'ZZTESTFROM' in v for v in row)]
    check(trips_cell and 3 in [v for v in trips_cell[0] if isinstance(v, (int, float))], f'the export counts all 3 trips ({trips_cell[:1]})')

    # the phone app and Dashboard keep their own period (not changed)
    r = cl.get('/app/analytics?tab=best')
    check(r.status_code == 200, 'the phone app Route Analytics page still opens')
finally:
    template_rendered.disconnect(grab, m.app)

db.close()
print(f'\nSUMMARY: {ok} passed, {fail} failed')
shutil.rmtree(tmp, ignore_errors=True)
print('Disposable copy removed. Real fleet.db was never modified by this test.')
sys.exit(1 if fail else 0)
