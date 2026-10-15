"""Tests for the Route Rates tab (one row per route: own -> party, hired -> party, hired -> owner, hired margin, own-vehicle driver advance / fuel).
Uses made-up trips on a throwaway database built from schema.sql, so every expected number is known. Also proves the numbers are LIVE: adding,
editing and deleting a trip changes the page on the next load, on the website and in the phone app, and the Excel export carries the same figures.

Usage:  python3 test_route_rates_layout.py
"""
import io
import os
import re
import shutil
import sqlite3
import sys
import tempfile

REPO = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, REPO)
os.chdir(REPO)
import requests                                                              # noqa: E402
def _blocked(*a, **k):
    raise RuntimeError('network blocked in tests')
requests.get = requests.post = _blocked
from openpyxl import load_workbook                                            # noqa: E402
import app as m                                                               # noqa: E402

ok = fail = 0


def check(cond, msg):
    global ok, fail
    if cond:
        ok += 1
    else:
        fail += 1
        print(f'[FAIL] {msg}')


tmp = tempfile.mkdtemp(prefix='rrlayout_')
dbp = os.path.join(tmp, 't.db')
c = sqlite3.connect(dbp)
c.executescript(open('schema.sql', encoding='utf-8').read())
c.execute("INSERT INTO users (id, username, password_hash, role, is_admin, full_name, access_level, status, phone, created_at) VALUES (1,'9000000001','x','Administrator',1,'Admin','Full Access','Active','9000000001','2026-01-01')")
c.execute("INSERT INTO vehicles (id, vehicle_no, type) VALUES (1,'OWN0001','own'),(2,'HIR0001','hired'),(3,'OWN0002','own')")
c.execute("INSERT INTO parties (id, name) VALUES (1,'ACME')")


def trip(vid, lr, frm, to, rate, qty=10, rtype='PER_MT', fuel=0, adv=0, owner_type=None, owner_amt=0, owner_rate=0, is_empty=0, fixed=0):
    vt = 'hired' if vid == 2 else 'own'
    c.execute("INSERT INTO trips (date, lr_number, vehicle_id, type, party_id, from_loc, to_loc, quantity, rate, rate_type, fixed_rate_amount, billed_amount, fuel_amount, "
              "driver_adv_amount, owner_rate_type, owner_fixed_amount, owner_rate, is_empty) VALUES ('2026-06-01',?,?,?,1,?,?,?,?,?,?,?,?,?,?,?,?,?)",
              (lr, vid, vt, frm, to, qty, rate, rtype, fixed, rate * qty, fuel, adv, owner_type or 'PER_MT', owner_amt, owner_rate, is_empty))


# route AAA -> BBB : own 1000 / 1200, hired 1500 / 1700 with fixed owner amounts 10000 / 12000 over 10 t (= 1000 / 1200 per tonne)
trip(1, 'O1', 'AAA', 'BBB', 1000, fuel=2000, adv=500)
trip(3, 'O2', 'AAA', 'BBB', 1200, fuel=3000, adv=0)
trip(2, 'H1', 'AAA', 'BBB', 1500, fuel=9999, adv=8888, owner_type='FIXED', owner_amt=10000)
trip(2, 'H2', 'AAA', 'BBB', 1700, owner_type='FIXED', owner_amt=12000)
# route CCC -> DDD : hired only; one trip with a per-tonne owner rate (800), one hired trip with no owner data
trip(2, 'H3', 'CCC', 'DDD', 1000, owner_type='PER_MT', owner_rate=800)
trip(2, 'H4', 'CCC', 'DDD', 900, owner_type='FIXED', owner_amt=0)
# route EEE -> FFF : hired trips with no owner data at all
trip(2, 'H5', 'EEE', 'FFF', 700)
# route GGG -> HHH : own only
trip(1, 'O3', 'GGG', 'HHH', 500, fuel=1000)
# never counted: a fixed-rate trip, an empty trip
trip(1, 'F1', 'AAA', 'BBB', 0, rtype='FIXED', fixed=50000)
trip(1, 'E1', 'AAA', 'BBB', 1, is_empty=1)
c.commit()
c.close()


def gdb():
    x = sqlite3.connect(dbp)
    x.row_factory = sqlite3.Row
    return x


m.get_db = gdb
m.app.config['TESTING'] = True
cl = m.app.test_client()
with cl.session_transaction() as s:
    s['user_id'] = 1


def rows():
    conn = gdb()
    d = m._route_rates_data(m._route_analytics_data(conn, None, None, '', '', '')['trips'], '', '')
    conn.close()
    return {(r['clean_from'], r['clean_to']): r for r in d['rr_rows']}, d


R, D = rows()
check(set(R) == {('AAA', 'BBB'), ('CCC', 'DDD'), ('EEE', 'FFF'), ('GGG', 'HHH')}, f'one row per route, fixed-rate and empty trips left out ({sorted(R)})')
a = R[('AAA', 'BBB')]
check((a['own']['highest'], a['own']['average'], a['own']['lowest'], a['own']['trips']) == (1200, 1100, 1000, 2), 'own -> party: 1200 / 1100 / 1000 over 2 trips')
check((a['hired_party']['highest'], a['hired_party']['average'], a['hired_party']['lowest'], a['hired_party']['trips']) == (1700, 1600, 1500, 2), 'hired -> party: 1700 / 1600 / 1500')
check((a['hired_owner']['highest'], a['hired_owner']['average'], a['hired_owner']['lowest'], a['hired_owner']['trips']) == (1200, 1100, 1000, 2), 'hired -> owner: fixed amounts per tonne (10000/10, 12000/10)')
check(a['margin_avg'] == 500 and round(a['margin_pct'], 2) == round(1000 / 3200 * 100, 2), f"hired margin = party - owner = 500 per tonne, {a['margin_pct']:.1f}% of the party rate")
check(a['op_trips'] == 2 and (a['da']['highest'], a['da']['trips']) == (500, 1) and (a['fuel']['highest'], a['fuel']['lowest'], a['fuel']['trips']) == (3000, 2000, 2),
      "own-vehicle details use OWN trips only (the hired trip's fuel 9999 / advance 8888 are not in them)")
b = R[('CCC', 'DDD')]
check(b['own'] is None and b['op_trips'] == 0 and b['da'] is None and b['fuel'] is None, 'hired-only route: no own columns, no operational details')
check(b['hired_party']['trips'] == 2 and b['hired_owner']['trips'] == 1 and b['hired_owner']['average'] == 800, 'a per-tonne owner rate is used as it is; a hired trip with no owner amount is left out of the owner figures')
check(b['margin_avg'] == 200 and round(b['margin_pct']) == 20, 'margin only over the trip that has both rates (1000 - 800 = 200, 20%)')
e = R[('EEE', 'FFF')]
check(e['hired_party'] and e['hired_owner'] is None and e['margin_avg'] is None and e['margin_pct'] is None, 'hired trips with no owner data: party rate shown, owner and margin empty')
g = R[('GGG', 'HHH')]
check(g['own']['trips'] == 1 and g['hired_party'] is None and g['hired_owner'] is None and g['margin_avg'] is None, 'own-only route: hired columns empty')
check(D['rr_total_routes'] == 4 and D['rr_total_count'] == 4, 'Total Routes counts distinct routes')

# ---------------- the website page
html = cl.get('/route-analytics?tab=route-rates').data.decode()
for text in ('Own Vehicle → Party (₹)', 'Hired Vehicle → Party (₹)', 'Hired Vehicle → Owner (₹)', 'Hired Margin (₹ / %)', 'No hired trips for this route', 'No owner rate data',
             'Own Vehicle Operational Details', 'These fields apply only to own vehicle trips'):
    check(text in html, f'website shows: {text}')
check(html.count('class="rr2-row"') == 4 and html.count('class="rr2-detail"') == 4, 'website: four routes, each with an expandable details row')
row = re.search(r'<tr class="rr2-row">.*?</tr>', html[html.index('AAA</td>') - 600:], re.S)
txt = re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', html[html.index('>AAA</td>') - 700: html.index('>AAA</td>') + 1800]))
check('₹1,200 ₹1,100 ₹1,000 2' in txt and '₹1,700 ₹1,600 ₹1,500 2' in txt and '₹500 31%' in txt, f'website row AAA → BBB carries the figures ({txt[txt.find("AAA"):][:200]})')
check('Hired Margin' in html and 'Driver Advances (₹)' in html and 'Fuel Taken (₹)' in html, 'the Driver Advances / Fuel Taken summary cards are still on the page')
check('name="from"' in html and 'name="vehicle_type"' in html and 'Highest Rate' in html, 'filters and top cards are untouched')

# ---------------- the phone app
ah = cl.get('/app/analytics?tab=rates').data.decode()
atxt = re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', ah))
check('AAA → BBB' in atxt and 'Own Vehicle → Party' in atxt and 'Hired Vehicle → Owner' in atxt and 'Hired Margin' in atxt, 'phone app shows the four blocks per route')
seg = atxt[atxt.index('AAA → BBB'): atxt.index('CCC → DDD')]
check('₹1,200' in seg and '₹1,100' in seg and '₹1,700' in seg and '₹500' in seg and '31%' in seg, 'phone app numbers match the website for AAA → BBB')
check('No hired vehicle data' not in atxt and 'No owner rate data' not in atxt and 'No own-vehicle trips' not in atxt and 'None recorded' not in atxt,
      'phone app: no "no data" messages anywhere — a block shows when it has data and is simply absent when it has none')
pc = {re.search(r'rr-row-route">([^<]+)<', ch).group(1): re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', ch)) for ch in (c.split('id="rr-load-area"')[0] for c in ah.split('class="rr-row"')[1:])}
hired_only, own_only, no_owner = pc['CCC → DDD'], pc['GGG → HHH'], pc['EEE → FFF']
check('Hired Vehicle → Party' in hired_only and 'Hired Vehicle → Owner' in hired_only and 'Own Vehicle → Party' not in hired_only and 'Own Vehicle Operational Details' not in hired_only,
      'phone app: a hired-only route has no own-vehicle block and no operational details')
check('Own Vehicle → Party' in own_only and 'Hired Vehicle' not in own_only and 'Hired Margin' not in own_only, 'phone app: an own-only route has no hired blocks at all')
check('Hired Vehicle → Party' in no_owner and 'Hired Vehicle → Owner' not in no_owner and 'Hired Margin' not in no_owner and 'Own Vehicle' not in no_owner,
      'phone app: hired trips without owner data show the party block only (no owner block, no margin)')
check('Own Vehicle Operational Details' in pc['AAA → BBB'] and 'Driver Advance' in pc['AAA → BBB'] and 'Fuel Taken' in pc['AAA → BBB'], 'phone app: a route with own trips shows its operational details')
check('Fuel Taken' in own_only and 'Driver Advance' not in own_only, 'phone app: the Driver Advance block is left out when no advance was given (only Fuel Taken shows)')
check('Own Vehicle Operational Details' in atxt and 'Driver Advance' in atxt and 'Fuel Taken' in atxt, 'phone app has the own-vehicle details')

# the phone list is compact: each route is one short row, and every value sits inside the part that opens when its arrow is tapped
chunks = ah.split('class="rr-row"')[1:]
check(len(chunks) == 4, f'phone app: four route rows ({len(chunks)})')
for ch in chunks:
    d, b = ch.find('class="rr-row-detail"'), ch.find('rr2m-block')
    route = re.search(r'rr-row-route">([^<]+)<', ch).group(1)
    check(d != -1 and b > d, f'phone app: {route} shows only its name until its arrow is tapped — the rate blocks are inside the expandable part')
    check('₹' not in ch[:d], f'phone app: {route} row itself carries no values')
check('class="rr-row-detail" id="rrDetail0"' in ah and 'display: none' in open('static/app.css', encoding='utf-8').read().split('.rr-row-detail {', 1)[1][:60], 'phone app: the expandable part is hidden until opened')

# ---------------- LIVE: the page follows the trips
def form(**kw):
    f = {'date': '2026-06-02', 'lr_number': 'NEW1', 'vehicle_no': 'HIR0001', 'party_name': 'ACME', 'from_loc': 'AAA', 'to_loc': 'BBB', 'rate_type': 'PER_MT', 'rate': '2000',
         'quantity': '10', 'owner_rate_type': 'FIXED', 'owner_fixed_amount': '14000'}
    f.update(kw)
    return f


r = cl.post('/trips/add', data=form())
check(r.status_code == 302, f'a hired trip added through the website ({r.status_code})')
R2, _ = rows()
a2 = R2[('AAA', 'BBB')]
check(a2['hired_party']['trips'] == 3 and a2['hired_party']['highest'] == 2000 and a2['hired_owner']['highest'] == 1400 and a2['hired_owner']['trips'] == 3,
      'adding a trip updates hired -> party and hired -> owner immediately')
check(round(a2['margin_avg'], 2) == round((500 + 500 + 600) / 3, 2), f"...and the margin ({a2['margin_avg']:.1f})")
html2 = cl.get('/route-analytics?tab=route-rates').data.decode()
check('₹2,000' in html2, 'the website page shows the new trip straight away')
check('₹2,000' in cl.get('/app/analytics?tab=rates').data.decode(), 'the phone app page shows it too')
nid = gdb().execute("SELECT id FROM trips WHERE lr_number='NEW1'").fetchone()[0]
ef = form(rate='2500', owner_fixed_amount='15000')
r = cl.post(f'/trips/edit/{nid}', data=ef)
R3, _ = rows()
check(R3[('AAA', 'BBB')]['hired_party']['highest'] == 2500 and R3[('AAA', 'BBB')]['hired_owner']['highest'] == 1500, f'editing a trip updates the rates immediately ({r.status_code})')
r = cl.post(f'/trips/delete/{nid}')
R4, _ = rows()
check(R4[('AAA', 'BBB')]['hired_party']['trips'] == 2 and R4[('AAA', 'BBB')]['hired_party']['highest'] == 1700, f'deleting a trip takes it out immediately ({r.status_code})')
r = cl.post('/app/trips/add', data=form(lr_number='NEW2', rate='1800', owner_fixed_amount='9000', from_loc='GGG', to_loc='HHH'))
R5, _ = rows()
check(R5[('GGG', 'HHH')]['hired_party'] and R5[('GGG', 'HHH')]['hired_party']['highest'] == 1800 and R5[('GGG', 'HHH')]['hired_owner']['highest'] == 900,
      f'a trip added in the phone app updates the same numbers ({r.status_code})')

# ---------------- filters and export
conn = gdb()
only_own = m._route_rates_data(m._route_analytics_data(conn, None, None, '', '', 'own')['trips'], '', '')['rr_rows']
check(all(r['hired_party'] is None for r in only_own) and any(r['own'] for r in only_own), 'the Vehicle Type = own filter leaves the hired columns empty')
check([(r['clean_from'], r['clean_to']) for r in m._route_rates_data(m._route_analytics_data(conn, None, None, 'AAA', 'BBB', '')['trips'], 'AAA', 'BBB')['rr_rows']] == [('AAA', 'BBB')], 'the From / To filters narrow to one route')
conn.close()
x = cl.get('/route-analytics/export')
wb = load_workbook(io.BytesIO(x.data))
ws = wb['Route Rates']
head = [cell.value for cell in ws[1]]
check('Hired to Owner: Average' in head and 'Hired Margin Avg (%)' in head and 'Own Fuel Average' in head, 'the Excel export has the new Route Rates columns')
line = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == 'AAA → BBB')
check(line[head.index('Own to Party: Average')] == 1100 and line[head.index('Hired to Owner: Average')] == 1100 and round(line[head.index('Hired Margin Avg (Rs/t)')]) == 500, 'the export row for AAA → BBB equals the page')

# ---------------- paging stays on its own tab and refreshes only its table
c2 = sqlite3.connect(dbp)
for i in range(1, 13):
    c2.execute("INSERT INTO trips (date, lr_number, vehicle_id, type, party_id, from_loc, to_loc, quantity, rate, rate_type, billed_amount) VALUES ('2026-06-05',?,1,'own',1,?,?,10,900,'PER_MT',9000)", (f'PG{i}', f'P{i:02d}', 'ZZZ'))
c2.commit(); c2.close()
pg1 = cl.get('/route-analytics').data.decode()      # how people really arrive: no tab in the address — the Route Rates tab is then opened by clicking its button in the browser
check('id="rrPager"' in pg1 and 'tab=route-rates' in re.search(r'id="rrPager".*?</div>\s*</div>', pg1, re.S).group(0), 'the Route Rates page links carry tab=route-rates (that is what keeps you on the tab)')
rr_links = re.findall(r'<a class="pager-btn pjax-page[^"]*" href="([^"]+)"', pg1[pg1.index('id="rrPager"'):])
check(rr_links and all('tab=route-rates' in h and 'rr_page=' in h for h in rr_links), f'every Route Rates pager link names its tab and its page ({len(rr_links)} links)')
best_links = re.findall(r'<a class="pager-btn pjax-page[^"]*" href="([^"]+)"', pg1[pg1.index('id="pager"'):pg1.index('id="ra-route-rates"')])
check(best_links and all('tab=best-routes' in h for h in best_links), 'the Best Routes pager links name their own tab too')
sel = re.search(r'<select data-pager-qs="([^"]*)" data-pager-anchor="rrPager"', pg1)
check(sel and 'tab=route-rates' in sel.group(1), 'the rows-per-page box carries the tab as well')
pg2 = cl.get('/route-analytics?tab=route-rates&rr_page=2').data.decode()
check(re.search(r'id="tabBtn-route-rates"[^>]*class="ra-tab active"|class="ra-tab active"[^>]*id="tabBtn-route-rates"', pg2) is not None, 'page 2 of Route Rates opens on the Route Rates tab (not Best Routes)')
rows2 = re.findall(r'<td class="sticky-col sticky-from"><button[^>]*>[^<]*</button>([^<]+)</td>', pg2)
rows1 = re.findall(r'<td class="sticky-col sticky-from"><button[^>]*>[^<]*</button>([^<]+)</td>', pg1)
check(len(rows1) == 10 and len(rows2) == 6 and not set(rows1) & set(rows2) and rows2[-1] == 'P12', f'page 1 has 10 routes, page 2 the other 6 ({len(rows1)} / {len(rows2)})')
check('window.__raPagerBound' in pg1 and 'data-pager-qs' in pg1, 'the page carries the refresh-only-the-table script')
soft = cl.get('/route-analytics?tab=route-rates&rr_page=2', headers={'X-Soft-Nav': '1'}).data.decode()
check('id="rrPager"' in soft and 'P12' in soft, 'the table-only refresh request returns the next page of rows')

# ---------------- design checks on the template
src = open('templates/route_analytics.html', encoding='utf-8').read()
check(all(k in src for k in ('--g-own', '--g-hp', '--g-ho', '--g-hm')) and 'data-theme="dark"] .rr-table' in src, 'four group colours, with brighter dark-theme values')

print(f'\nSUMMARY: {ok} passed, {fail} failed')
shutil.rmtree(tmp, ignore_errors=True)
sys.exit(1 if fail else 0)
